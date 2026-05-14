"""Build McKinsey/MBB-style pitch deck for EBS Agentic Support Agent."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).parent
SHOTS = ROOT / "screens"
OUT = ROOT / "EBS_Agentic_Support_Demo.pptx"

# McKinsey palette
NAVY     = RGBColor(0x05, 0x1C, 0x2C)   # deep navy
BLUE     = RGBColor(0x22, 0x51, 0xFF)   # accent blue
MID      = RGBColor(0x03, 0x4B, 0x73)
GREY     = RGBColor(0x8C, 0x9C, 0xAA)
LIGHT    = RGBColor(0xE6, 0xE6, 0xE6)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
BLACK    = RGBColor(0x00, 0x00, 0x00)

FONT = "Georgia"      # McKinsey uses Bower / serif-ish for headers; Georgia is a safe local stand-in
FONT_SANS = "Arial"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, *, size=14, color=NAVY, bold=False,
             font=FONT_SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def chrome(slide, page_num, section="", title_kicker="EBS AGENTIC SUPPORT"):
    # top hairline
    add_rect(slide, Inches(0.5), Inches(0.45), Inches(12.33), Emu(9525), NAVY)
    # tiny kicker left
    add_text(slide, Inches(0.5), Inches(0.15), Inches(8), Inches(0.25),
             title_kicker, size=8, color=GREY, bold=True, font=FONT_SANS)
    # section right
    if section:
        add_text(slide, Inches(8), Inches(0.15), Inches(4.83), Inches(0.25),
                 section, size=8, color=GREY, bold=True, font=FONT_SANS, align=PP_ALIGN.RIGHT)
    # bottom hairline
    add_rect(slide, Inches(0.5), Inches(7.1), Inches(12.33), Emu(9525), LIGHT)
    # footer
    add_text(slide, Inches(0.5), Inches(7.2), Inches(8), Inches(0.22),
             "CONFIDENTIAL  |  CLIENT DISCUSSION DOCUMENT", size=7.5, color=GREY, font=FONT_SANS)
    add_text(slide, Inches(11.8), Inches(7.2), Inches(1.03), Inches(0.22),
             str(page_num), size=7.5, color=GREY, font=FONT_SANS, align=PP_ALIGN.RIGHT)


# ---------------- Slide 1: cover ----------------
s = prs.slides.add_slide(BLANK)
# full navy panel left
add_rect(s, 0, 0, Inches(5), SH, NAVY)
# accent vertical line
add_rect(s, Inches(5), 0, Emu(19050), SH, BLUE)
# kicker
add_text(s, Inches(0.6), Inches(0.7), Inches(4), Inches(0.3),
         "CLIENT BRIEFING", size=9, color=GREY, bold=True)
# big title on right
add_text(s, Inches(5.6), Inches(2.1), Inches(7.3), Inches(1.2),
         "Oracle EBS", size=46, color=NAVY, bold=False, font=FONT)
add_text(s, Inches(5.6), Inches(2.9), Inches(7.3), Inches(1.2),
         "Agentic Support", size=46, color=BLUE, bold=False, font=FONT)
add_text(s, Inches(5.6), Inches(4.1), Inches(7.3), Inches(0.5),
         "An AI-native operations copilot for the EBS estate", size=15, color=GREY, font=FONT_SANS)
# bottom meta
add_rect(s, Inches(5.6), Inches(6.1), Inches(1), Emu(15875), NAVY)
add_text(s, Inches(5.6), Inches(6.25), Inches(7), Inches(0.4),
         "Prepared for   ·   Big Four Advisory", size=10, color=NAVY, bold=True)
add_text(s, Inches(5.6), Inches(6.55), Inches(7), Inches(0.4),
         "April 2026", size=9, color=GREY)
# left-panel vertical label
add_text(s, Inches(0.6), Inches(6.6), Inches(4), Inches(0.3),
         "DEMONSTRATION", size=9, color=GREY, bold=True)


# ---------------- Slide 2: context ----------------
s = prs.slides.add_slide(BLANK); chrome(s, 2, "01  ·  CONTEXT")
add_text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
         "The EBS operations gap", size=28, color=NAVY, font=FONT)
add_text(s, Inches(0.5), Inches(1.55), Inches(11), Inches(0.4),
         "Enterprises run EBS for decades — the support model has not kept up.",
         size=13, color=GREY, font=FONT_SANS)

# three pillars
cols = [
    ("Fragmented", "Seventeen modules,\nhundreds of diagnostics,\nno single pane."),
    ("Reactive", "Issues surface in\nmonth-end close —\nnot before."),
    ("Expensive", "Tier-1 support\ntickets crowd out\nstrategic work."),
]
x0 = Inches(0.5); y0 = Inches(3.2); cw = Inches(4.0); gap = Inches(0.15)
for i, (h, b) in enumerate(cols):
    x = x0 + (cw + gap) * i
    add_rect(s, x, y0, Emu(38100), Inches(0.02), BLUE)
    add_text(s, x, y0 + Inches(0.18), cw, Inches(0.5), h,
             size=18, color=NAVY, bold=True, font=FONT)
    add_text(s, x, y0 + Inches(0.9), cw, Inches(2.5), b,
             size=12, color=BLACK, font=FONT_SANS, spacing=1.3)


# ---------------- Slide 3: solution stats ----------------
s = prs.slides.add_slide(BLANK); chrome(s, 3, "02  ·  SOLUTION")
add_text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
         "One agent. The full estate.", size=28, color=NAVY, font=FONT)
add_text(s, Inches(0.5), Inches(1.55), Inches(11), Inches(0.4),
         "Conversational, live, grounded in the customer's own EBS.",
         size=13, color=GREY, font=FONT_SANS)

stats = [("117", "diagnostic agents"),
         ("15",  "agentic applications"),
         ("30+", "live Oracle queries"),
         ("3",   "functional pillars")]
x0 = Inches(0.5); y0 = Inches(3.3); cw = Inches(3.0); gap = Inches(0.08)
for i, (n, lab) in enumerate(stats):
    x = x0 + (cw + gap) * i
    add_text(s, x, y0, cw, Inches(1.4), n, size=72, color=BLUE, font=FONT, bold=False)
    add_rect(s, x, y0 + Inches(1.5), Inches(0.4), Emu(15875), NAVY)
    add_text(s, x, y0 + Inches(1.6), cw, Inches(0.5), lab,
             size=11, color=NAVY, bold=True, font=FONT_SANS)

add_text(s, Inches(0.5), Inches(6.3), Inches(12), Inches(0.4),
         "Finance  ·  Supply Chain  ·  Planning",
         size=11, color=GREY, bold=True, font=FONT_SANS)


# ---------------- Slide 4: architecture ----------------
s = prs.slides.add_slide(BLANK); chrome(s, 4, "03  ·  HOW IT WORKS")
add_text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
         "Architecture at a glance", size=28, color=NAVY, font=FONT)

layers = [
    ("Interface",   "Chat  ·  Dashboards  ·  Agentic Apps"),
    ("Orchestrator","Claude Sonnet 4.6  ·  Multi-agent routing"),
    ("Agents",      "117 diagnostic  ·  15 operational"),
    ("Data",        "Live Oracle EBS  ·  PostgreSQL audit"),
]
x = Inches(0.5); y = Inches(2.4); w = Inches(12.33); h = Inches(0.95)
for i, (a, b) in enumerate(layers):
    yy = y + (h + Inches(0.08)) * i
    add_rect(s, x, yy, Inches(0.08), h, BLUE if i == 0 else NAVY)
    add_text(s, x + Inches(0.3), yy + Inches(0.1), Inches(3), Inches(0.4),
             a.upper(), size=10, color=GREY, bold=True, font=FONT_SANS)
    add_text(s, x + Inches(0.3), yy + Inches(0.4), Inches(11.5), Inches(0.5),
             b, size=16, color=NAVY, font=FONT)


# ---------------- Screenshot slides ----------------
def shot_slide(page, section, headline, subhead, img_name, *, crop=False):
    s = prs.slides.add_slide(BLANK); chrome(s, page, section)
    add_text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
             headline, size=26, color=NAVY, font=FONT)
    add_text(s, Inches(0.5), Inches(1.5), Inches(11.5), Inches(0.4),
             subhead, size=12, color=GREY, font=FONT_SANS)
    # image frame — subtle border
    img_path = SHOTS / img_name
    # place image fit-to-width
    frame_x = Inches(0.5); frame_y = Inches(2.1)
    frame_w = Inches(12.33); frame_h = Inches(4.85)
    # thin accent bar on top-left of the image
    add_rect(s, frame_x, frame_y - Inches(0.05), Inches(0.5), Emu(19050), BLUE)
    # background
    add_rect(s, frame_x, frame_y, frame_w, frame_h, LIGHT)
    # add picture centered inside
    pic = s.shapes.add_picture(str(img_path), frame_x, frame_y, width=frame_w)
    # if too tall, re-insert height-bound
    if pic.height > frame_h:
        sp = pic._element; sp.getparent().remove(sp)
        pic = s.shapes.add_picture(str(img_path), frame_x, frame_y, height=frame_h)
        pic.left = frame_x + (frame_w - pic.width) // 2
    else:
        pic.top = frame_y + (frame_h - pic.height) // 2
        pic.left = frame_x + (frame_w - pic.width) // 2


shot_slide(5, "04  ·  DEMO",
           "Conversational diagnostics",
           "Ask in plain English — the orchestrator routes, queries EBS, returns grounded answers.",
           "chat.png")

shot_slide(6, "04  ·  DEMO",
           "117-agent catalog",
           "Every diagnostic discoverable, filterable, and callable from a single surface.",
           "dashboard.png")

shot_slide(7, "04  ·  DEMO",
           "EBS Agentic Apps",
           "Fifteen operational apps across Finance, Supply Chain, and Planning.",
           "fusion.png")

shot_slide(8, "04  ·  DEMO",
           "Live execution, drill to row",
           "One click runs the app against live Oracle EBS and opens the exception drawer.",
           "fusion_drawer.png")

shot_slide(9, "04  ·  DEMO",
           "Supply Chain Planning",
           "Six KPIs, seventeen views, row-level recommendations tied to live plan data.",
           "scp.png")

shot_slide(10, "04  ·  DEMO",
           "Planning exceptions — actioned",
           "Shortages, excess, late supply, safety-stock breaches — surfaced with next best action.",
           "scp_results.png")

shot_slide(11, "04  ·  DEMO",
           "Payables copilot",
           "Discrepancies, holds, and cash-impact views ready for the controller.",
           "payables.png")


# ---------------- Slide 12: value ----------------
s = prs.slides.add_slide(BLANK); chrome(s, 12, "05  ·  VALUE")
add_text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
         "Where value shows up", size=28, color=NAVY, font=FONT)

rows = [
    ("Support cost",    "Tier-1 ticket volume",     "−40 to 60%"),
    ("Close cycle",     "Days to variance insight", "−3 to 5 days"),
    ("Working capital", "Unapplied receipts AR",    "−15 to 25%"),
    ("Plan quality",    "Forecast MAPE",            "−5 to 10 pts"),
]
y = Inches(2.4)
for i, (a, b, c) in enumerate(rows):
    yy = y + Inches(0.85) * i
    add_rect(s, Inches(0.5), yy + Inches(0.75), Inches(12.33), Emu(9525), LIGHT)
    add_text(s, Inches(0.5), yy, Inches(3), Inches(0.5), a,
             size=11, color=GREY, bold=True, font=FONT_SANS)
    add_text(s, Inches(3.5), yy, Inches(5), Inches(0.5), b,
             size=16, color=NAVY, font=FONT)
    add_text(s, Inches(9.0), yy, Inches(3.83), Inches(0.5), c,
             size=20, color=BLUE, bold=True, font=FONT, align=PP_ALIGN.RIGHT)

add_text(s, Inches(0.5), Inches(6.5), Inches(12), Inches(0.3),
         "Indicative — to be validated in pilot.", size=9, color=GREY, font=FONT_SANS)


# ---------------- Slide 13: engagement ----------------
s = prs.slides.add_slide(BLANK); chrome(s, 13, "06  ·  NEXT")
add_text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
         "A six-week pilot", size=28, color=NAVY, font=FONT)

phases = [
    ("01", "Connect",  "Week 1",    "Secure read-only link to EBS. One sandbox."),
    ("02", "Configure","Week 2-3",  "Tune agents to client schema and playbooks."),
    ("03", "Co-pilot", "Week 4-5",  "Shadow live tickets with measured deflection."),
    ("04", "Decide",   "Week 6",    "Business case, scope for scale."),
]
x0 = Inches(0.5); y0 = Inches(2.6); cw = Inches(3.0); gap = Inches(0.08)
for i, (n, h, w_, b) in enumerate(phases):
    x = x0 + (cw + gap) * i
    add_text(s, x, y0, cw, Inches(0.6), n,
             size=32, color=BLUE, font=FONT)
    add_rect(s, x, y0 + Inches(0.7), Inches(0.6), Emu(15875), NAVY)
    add_text(s, x, y0 + Inches(0.85), cw, Inches(0.4), h.upper(),
             size=12, color=NAVY, bold=True, font=FONT_SANS)
    add_text(s, x, y0 + Inches(1.25), cw, Inches(0.3), w_,
             size=9, color=GREY, bold=True, font=FONT_SANS)
    add_text(s, x, y0 + Inches(1.7), cw, Inches(2), b,
             size=11, color=BLACK, font=FONT_SANS, spacing=1.3)


# ---------------- Slide 14: closing ----------------
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, NAVY)
add_rect(s, Inches(0.7), Inches(0.7), Inches(0.4), Emu(19050), BLUE)
add_text(s, Inches(0.7), Inches(0.9), Inches(6), Inches(0.3),
         "THANK YOU", size=10, color=GREY, bold=True, font=FONT_SANS)
add_text(s, Inches(0.7), Inches(2.6), Inches(12), Inches(1.5),
         "Let's put it in front", size=48, color=WHITE, font=FONT)
add_text(s, Inches(0.7), Inches(3.5), Inches(12), Inches(1.5),
         "of real tickets.", size=48, color=BLUE, font=FONT)
add_text(s, Inches(0.7), Inches(5.2), Inches(6), Inches(0.4),
         "Discussion", size=10, color=GREY, bold=True, font=FONT_SANS)
add_text(s, Inches(0.7), Inches(5.5), Inches(12), Inches(0.5),
         "Questions  ·  Architecture deep-dive  ·  Pilot scoping",
         size=16, color=WHITE, font=FONT)


# ---------------- Appendix: 117-agent catalog ----------------
import sys
sys.path.insert(0, str(ROOT.parent))
import config.analyzer_registry as _reg
_agents = list(_reg.ANALYZER_REGISTRY.values())
# group by module
from collections import defaultdict
by_mod = defaultdict(list)
for a in _agents:
    by_mod[a.get("module", "OTHER")].append(a)
# stable ordering
_mod_order = ["ATG", "FINANCIALS", "MANUFACTURING", "HCM", "CRM"]
_mods = [m for m in _mod_order if m in by_mod] + [m for m in sorted(by_mod) if m not in _mod_order]

# Appendix divider
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, NAVY)
add_rect(s, Inches(0.7), Inches(0.7), Inches(0.4), Emu(19050), BLUE)
add_text(s, Inches(0.7), Inches(0.9), Inches(6), Inches(0.3),
         "APPENDIX", size=10, color=GREY, bold=True, font=FONT_SANS)
add_text(s, Inches(0.7), Inches(2.8), Inches(12), Inches(1.2),
         "The 117 agents.", size=56, color=WHITE, font=FONT)
add_text(s, Inches(0.7), Inches(4.0), Inches(12), Inches(0.5),
         f"Grouped across {len(_mods)} EBS functional modules.",
         size=16, color=GREY, font=FONT_SANS)

# Agent list slides — paginated, columnar
PAGE_ROWS = 22
PAGE_COLS = 3

def _agent_pages():
    """Yield (module, [agents chunk]) tuples, chunked per page."""
    for mod in _mods:
        items = sorted(by_mod[mod], key=lambda x: x["name"].lower())
        per_page = PAGE_ROWS * PAGE_COLS
        for i in range(0, len(items), per_page):
            yield mod, items[i:i + per_page], i, len(items)

_page_num = 15
for mod, chunk, offset, total in _agent_pages():
    s = prs.slides.add_slide(BLANK)
    chrome(s, _page_num, f"APPENDIX  ·  {mod}")
    _page_num += 1
    add_text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.5),
             f"{mod.title()} agents", size=26, color=NAVY, font=FONT)
    rng_hi = min(offset + len(chunk), total)
    add_text(s, Inches(0.5), Inches(1.5), Inches(12), Inches(0.4),
             f"{offset+1}–{rng_hi} of {total}", size=11, color=GREY, font=FONT_SANS)

    col_w = Inches(4.05); col_gap = Inches(0.1)
    y0 = Inches(2.1)
    row_h = Inches(0.22)
    for idx, a in enumerate(chunk):
        col = idx // PAGE_ROWS
        row = idx % PAGE_ROWS
        x = Inches(0.5) + col * (col_w + col_gap)
        y = y0 + row * row_h
        # bullet dot
        add_rect(s, x, y + Emu(45720), Emu(38100), Emu(38100), BLUE)
        name = a["name"]
        if len(name) > 46:
            name = name[:44] + "…"
        add_text(s, x + Inches(0.12), y, col_w - Inches(0.12), row_h,
                 name, size=9, color=NAVY, font=FONT_SANS)

prs.save(str(OUT))
print(f"wrote {OUT}  ·  {len(prs.slides)} slides")
